#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""PPT 视觉审稿质检（slide-designer 交付前自检脚本）

借鉴来源：@yejiming/dsh-ppt（AI 视觉审稿）、liustack/pptwise（实时预览迭代）、
dream-num/dsh-univer-office（文本溢出/重叠检测）。

用法:
    python visual-qa.py <deck.pptx>

退出码:
    0 = 通过（无确定性告警）
    1 = 发现问题（按报告逐条处理，或人工确认后放行）
    2 = 缺依赖 / 文件不存在 / 参数错误

说明:
    - 只做「几何 + 估算」两类确定性检查：越界 / 重叠 / 文字溢出估算 / 对比度 / 字体。
    - 文字溢出为「估算」：全角字宽≈字号、ASCII≈0.52×字号、行高≈1.25×字号。
      最终以 LibreOffice / PowerPoint 渲染缩略图人工复核为准（见 knowledge/visual-qa.md）。
    - 依赖 python-pptx（pip install python-pptx）。缺依赖时按 visual-qa.md 人工目检清单执行。
"""
import sys
import os
import math

PT_EMU = 12700.0  # 1 pt = 12700 EMU


def rel_lum(rgb):
    def ch(c):
        c = c / 255.0
        return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4
    return 0.2126 * ch(rgb[0]) + 0.7152 * ch(rgb[1]) + 0.0722 * ch(rgb[2])


def contrast(a, b):
    la, lb = rel_lum(a), rel_lum(b)
    hi, lo = max(la, lb), min(la, lb)
    return (hi + 0.05) / (lo + 0.05)


def run_rgb(run):
    try:
        c = run.font.color
        if c is None or getattr(c, "type", None) is None or c.rgb is None:
            return None
        return (c.rgb[0], c.rgb[1], c.rgb[2])
    except Exception:
        return None


def shape_fill_rgb(shape):
    try:
        f = shape.fill
        if f is None or getattr(f, "type", None) is None:
            return None
        if "SOLID" in str(f.type).upper():
            r = f.fore_color.rgb
            return (r[0], r[1], r[2])
    except Exception:
        pass
    return None


def char_width(cp, size):
    # 中日韩 + 全角区（含全角标点）≈ 1.0×字号；其余（ASCII/半角）≈ 0.52×字号
    if cp >= 0x2E80 or 0x3000 <= cp <= 0x303F or 0xFF00 <= cp <= 0xFFEF or 0x2000 <= cp <= 0x206F:
        return size
    return 0.52 * size


def text_width_pt(text, size):
    return sum(char_width(ord(c), size) for c in text)


def check_overflow(shape):
    if not shape.has_text_frame:
        return None
    text = shape.text_frame.text
    if not text.strip():
        return None
    w = shape.width / PT_EMU
    h = shape.height / PT_EMU
    size = 18.0
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            if r.font.size is not None:
                size = max(size, r.font.size.pt)
    flat = text.replace("\n", "").replace("\v", "").replace("\r", "")
    total_w = text_width_pt(flat, size)
    hard_lines = text.count("\n") + text.count("\v") + 1
    wrap_lines = max(1, int(math.ceil(total_w / max(w, 1.0))))
    lines = max(hard_lines, wrap_lines)
    needed = lines * size * 1.25
    if needed > h * 1.05 + 2:
        return "文字溢出(估算) 约需 %.0fpt 高 / 框高 %.0fpt / 预估 %d 行" % (needed, h, lines)
    return None


def intersect_area(a, b):
    l1, t1, r1, b1 = a
    l2, t2, r2, b2 = b
    w = min(r1, r2) - max(l1, l2)
    h = min(b1, b2) - max(t1, t2)
    return max(0.0, w) * max(0.0, h)


def area(a):
    l, t, r, b = a
    return max(0.0, r - l) * max(0.0, b - t)


def main():
    if len(sys.argv) < 2:
        print("用法: python visual-qa.py <deck.pptx>")
        return 2
    path = sys.argv[1]
    if not os.path.exists(path):
        print("文件不存在: " + path)
        return 2
    try:
        from pptx import Presentation
    except ImportError:
        print("[缺依赖] 未安装 python-pptx。执行 `pip install python-pptx` 后重跑；")
        print("        或按 knowledge/visual-qa.md「人工目检清单」逐项检查。")
        return 2

    SAFE_FONTS = {
        "microsoft yahei", "微软雅黑", "simhei", "黑体", "simsun", "宋体",
        "dengxian", "等线", "arial", "calibri", "segoe ui", "times new roman",
        "helvetica", "inter", "source han sans", "思源黑体", "思源宋体", "jetbrains mono",
    }

    prs = Presentation(path)
    sw, sh = prs.slide_width / PT_EMU, prs.slide_height / PT_EMU
    total = 0
    for i, slide in enumerate(prs.slides, 1):
        shapes = list(slide.shapes)
        boxes = []  # (shape, l, t, r, b)
        for s in shapes:
            try:
                if s.left is None or s.top is None or s.width is None or s.height is None:
                    continue
                l = s.left / PT_EMU
                t = s.top / PT_EMU
                r = l + s.width / PT_EMU
                b = t + s.height / PT_EMU
                boxes.append((s, l, t, r, b))
            except Exception:
                continue

        # 1) 越界
        for s, l, t, r, b in boxes:
            if l < -2 or t < -2 or r > sw + 2 or b > sh + 2:
                print("[越界] slide %d: '%s' 超出页面 (%.0f,%.0f)-(%.0f,%.0f), 页面 %.0f×%.0f"
                      % (i, (s.name or "?")[:30], l, t, r, b, sw, sh))
                total += 1

        # 2) 文字溢出 + 对比度 + 字体
        for s, l, t, r, b in boxes:
            if s.has_text_frame:
                ov = check_overflow(s)
                if ov:
                    print("[%s] slide %d: '%s'" % (ov, i, (s.name or "?")[:30]))
                    total += 1
                fill = shape_fill_rgb(s)
                for p in s.text_frame.paragraphs:
                    for run in p.runs:
                        fg = run_rgb(run)
                        if fg and fill:
                            c = contrast(fg, fill)
                            if c < 2.5:
                                print("[对比度不足 %.1f:1] slide %d: '%s' 文本色%s vs 底色%s"
                                      % (c, i, (s.name or "?")[:30], fg, fill))
                                total += 1
                        nm = (run.font.name or "").lower()
                        if nm and nm not in SAFE_FONTS:
                            print("[字体需确认] slide %d: '%s' 使用 '%s'，确认目标机可用/已嵌入"
                                  % (i, (s.name or "?")[:30], run.font.name))
                            total += 1

        # 3) 文本形状两两重叠
        text_boxes = [(s, l, t, r, b) for s, l, t, r, b in boxes
                      if s.has_text_frame and s.text_frame.text.strip()]
        for a in range(len(text_boxes)):
            for b_idx in range(a + 1, len(text_boxes)):
                sa, la, ta, ra, ba = text_boxes[a]
                sb, lb, tb, rb, bb = text_boxes[b_idx]
                ia = intersect_area((la, ta, ra, ba), (lb, tb, rb, bb))
                small = min(area((la, ta, ra, ba)), area((lb, tb, rb, bb)))
                if small > 0 and ia / small > 0.15:
                    print("[重叠需确认] slide %d: '%s' 与 '%s' 相交 %.0f%%"
                          % (i, (sa.name or "?")[:24], (sb.name or "?")[:24], 100 * ia / small))
                    total += 1

    print("\n===== 视觉审稿完成：%d 项告警 =====" % total)
    print("提示：溢出为估算值；重叠/对比度/字体需人工结合渲染缩略图最终确认（见 knowledge/visual-qa.md）。")
    return 1 if total else 0


if __name__ == "__main__":
    sys.exit(main())
